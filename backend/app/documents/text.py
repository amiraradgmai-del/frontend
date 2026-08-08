from __future__ import annotations

import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass

import pymupdf
from docx import Document as WordDocument
from openpyxl import load_workbook

_CHARACTER_TRANSLATION = str.maketrans(
    {"ي": "ی", "ى": "ی", "ك": "ک", "ة": "ه", "ۀ": "ه"}
)
_ARABIC_DIACRITICS = re.compile(r"[\u064b-\u065f\u0670\u06d6-\u06ed]")
_ARTICLE_PATTERN = re.compile(
    r"(?<!\S)ماده\s+([۰-۹0-9]+(?:\s+مکرر)?)\s*[-–—:]?",
    re.MULTILINE,
)


class ExtractionError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True, slots=True)
class PageText:
    page_number: int | None
    text: str


@dataclass(frozen=True, slots=True)
class ChunkDraft:
    chunk_index: int
    content: str
    content_hash: str
    page_number: int | None
    article_number: str | None
    section_title: str | None
    char_count: int
    token_count: int


def normalize_persian(text: str, *, preserve_lines: bool = True) -> str:
    text = unicodedata.normalize("NFKC", text).translate(_CHARACTER_TRANSLATION)
    text = _ARABIC_DIACRITICS.sub("", text).replace("\ufeff", "").replace("\u200f", "")
    text = re.sub(r"[\t\r\f\v]+", " ", text)
    text = re.sub(r"[ \u00a0]+", " ", text)
    if preserve_lines:
        lines = [line.strip() for line in text.splitlines()]
        return "\n".join(line for line in lines if line)
    return re.sub(r"\s+", " ", text).strip()


def extract_pages(data: bytes, extension: str) -> list[PageText]:
    try:
        if extension == ".pdf":
            with pymupdf.open(stream=data, filetype="pdf") as document:
                return [
                    PageText(index + 1, normalize_persian(page.get_text("text")))
                    for index, page in enumerate(document)
                ]
        if extension == ".docx":
            document = WordDocument(io.BytesIO(data))
            blocks = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                blocks.extend(
                    " | ".join(cell.text for cell in row.cells) for row in table.rows
                )
            return [PageText(None, normalize_persian("\n".join(blocks)))]
        if extension == ".xlsx":
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            pages: list[PageText] = []
            try:
                for index, sheet in enumerate(workbook.worksheets, 1):
                    rows = []
                    for row in sheet.iter_rows(values_only=True):
                        values = [str(value).strip() for value in row if value not in (None, "")]
                        if values:
                            rows.append(" | ".join(values))
                    pages.append(PageText(index, normalize_persian(f"برگه: {sheet.title}\n" + "\n".join(rows))))
            finally:
                workbook.close()
            return pages
        if extension in {".jpg", ".jpeg", ".png", ".webp"}:
            return [PageText(1, "")]
        if extension == ".txt":
            return [PageText(None, normalize_persian(data.decode("utf-8-sig")))]
    except Exception as error:
        if isinstance(error, ExtractionError):
            raise
        raise ExtractionError(
            "extraction_failed", "Document text extraction failed"
        ) from error
    raise ExtractionError("unsupported_type", "Unsupported document type")


def _split_bounded(text: str, max_chars: int, overlap_chars: int) -> list[str]:
    text = normalize_persian(text, preserve_lines=False)
    if len(text) <= max_chars:
        return [text] if text else []
    chunks: list[str] = []
    start = 0
    while start < len(text):
        hard_end = min(start + max_chars, len(text))
        end = hard_end
        if hard_end < len(text):
            search_start = start + max_chars // 2
            candidates = [
                text.rfind(mark, search_start, hard_end)
                for mark in (". ", "؟ ", "؛ ", " ")
            ]
            boundary = max(candidates)
            if boundary >= search_start:
                end = boundary + 1
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = max(start + 1, end - overlap_chars)
    return chunks


def build_chunks(
    pages: list[PageText], *, max_chars: int, overlap_chars: int
) -> list[ChunkDraft]:
    drafts: list[ChunkDraft] = []
    for page in pages:
        matches = list(_ARTICLE_PATTERN.finditer(page.text))
        sections: list[tuple[str, str | None]] = []
        if not matches:
            sections.append((page.text, None))
        else:
            prefix = page.text[: matches[0].start()].strip()
            if prefix:
                sections.append((prefix, None))
            for index, match in enumerate(matches):
                end = (
                    matches[index + 1].start()
                    if index + 1 < len(matches)
                    else len(page.text)
                )
                sections.append((page.text[match.start() : end], match.group(1)))
        for section, article_number in sections:
            for content in _split_bounded(section, max_chars, overlap_chars):
                drafts.append(
                    ChunkDraft(
                        chunk_index=len(drafts),
                        content=content,
                        content_hash=hashlib.sha256(
                            content.encode("utf-8")
                        ).hexdigest(),
                        page_number=page.page_number,
                        article_number=article_number,
                        section_title=None,
                        char_count=len(content),
                        token_count=len(content.split()),
                    )
                )
    return drafts
