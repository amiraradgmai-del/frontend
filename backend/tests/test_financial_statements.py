from decimal import Decimal
from io import BytesIO

import fitz
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session

from app.db.base import Base
from app.models.financial_statements import *
from app.services.financial_statements import (
    AccountMappingService, ExcelTrialBalanceParser, LegacyExcelTrialBalanceParser,
    FinancialStatementCalculationService, convert_money, extract_account,
    normalize_persian_financial, normalized_account_name, parse_amount,
    build_export_data, pdf_export, workbook_export,
)
from app.schemas.financial_statements import FiscalYearCreate, OrganizationCreate


def excel_bytes() -> bytes:
    workbook = Workbook(); sheet = workbook.active
    sheet.append(["گزارش تراز شش ستونی"])
    sheet.append(["کل", "معین", "تفصیلی", "بدهکار اول دوره", "بستانکار اول دوره", "بدهکار طی دوره", "بستانکار طی دوره", "مانده بدهکار", "مانده بستانکار"])
    sheet.append(["1110 - موجودی نقد و بانک", "111003 - تنخواه گردان", "", 0, 0, 10_000_000, 0, 10_000_000, 0])
    sheet.append(["", "111005 - بانک ریالی", "بانک ملت", 0, 0, 50_000_000, 0, 50_000_000, 0])
    output = BytesIO(); workbook.save(output); return output.getvalue()


def test_persian_normalization_account_and_amount():
    assert normalized_account_name(" موجودي نقد‌ وبانك ") == "موجودی نقد وبانک"
    assert extract_account("۱۱۱۰ — موجودی نقد و بانک") == ("1110", "موجودی نقد و بانک")
    assert parse_amount("(۱٬۲۳۴٬۵۶۷ ریال)") == Decimal("-1234567")


def test_excel_parser_detects_variable_header_position():
    rows = ExcelTrialBalanceParser().parse(excel_bytes())
    assert len(rows) == 2
    assert rows[0].general_code == "1110"
    assert rows[1].subsidiary_code == "111005"
    assert rows[1].closing_debit == Decimal("50000000")


def test_fastreport_spreadsheet_xml_with_xls_extension():
    data = """<?xml version="1.0"?>
    <Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet"
      xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">
      <Worksheet ss:Name="Page 1"><Table>
        <Row><Cell ss:Index="2"><Data ss:Type="String">مانده بد</Data></Cell><Cell><Data ss:Type="String">بس طی دوره</Data></Cell><Cell><Data ss:Type="String">بد طی دوره</Data></Cell><Cell><Data ss:Type="String">بس اول دوره</Data></Cell><Cell ss:Index="7"><Data ss:Type="String">بد اول دوره</Data></Cell><Cell ss:Index="9"><Data ss:Type="String">تفصیلی</Data></Cell><Cell><Data ss:Type="String">معین</Data></Cell><Cell ss:Index="13"><Data ss:Type="String">کـل</Data></Cell></Row>
        <Row><Cell ss:Index="2"><Data ss:Type="Number">1000</Data></Cell><Cell><Data ss:Type="Number">0</Data></Cell><Cell><Data ss:Type="Number">0</Data></Cell><Cell><Data ss:Type="Number">0</Data></Cell><Cell ss:Index="7"><Data ss:Type="Number">1000</Data></Cell><Cell ss:Index="9"><Data ss:Type="String">صندوق</Data></Cell><Cell><Data ss:Type="String">111003 - تنخواه</Data></Cell><Cell ss:Index="13"><Data ss:Type="String">1110 - موجودی نقد و بانک</Data></Cell></Row>
      </Table></Worksheet>
    </Workbook>""".encode()
    rows = LegacyExcelTrialBalanceParser().parse(data)
    assert len(rows) == 1
    assert rows[0].general_code == "1110"
    assert rows[0].closing_debit == Decimal("1000")
    assert AccountMappingService._fallback(rows[0].general_code) == ("BS.CASH", 1)


def test_mapping_aggregation_sign_unit_validation_and_versioning(tmp_path):
    engine = create_engine(f"sqlite+pysqlite:///{tmp_path / 'financial.db'}")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        session.add(FinancialStatementField(id="BS.CASH", statement="BS", title="موجودی نقد", field_type="mapped"))
        organization = FinancialOrganization(name="شرکت آزمون", owner_user_id="user-1")
        session.add(organization); session.flush()
        year = FinancialFiscalYear(organization_id=organization.id, title="۱۴۰۵")
        session.add(year); session.flush()
        import_ = FinancialStatementImport(organization_id=organization.id, fiscal_year_id=year.id, user_id="user-1", original_filename="sample.xlsx", storage_key="sample", mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", file_size=10, file_hash="a" * 64, money_unit="rial", status="parsed")
        session.add(import_); session.flush()
        for index, amount in enumerate((Decimal("10000000"), Decimal("50000000")), 1):
            session.add(FinancialStatementImportRow(import_id=import_.id, general_code="1110", general_name="موجودی نقد و بانک", subsidiary_code=f"11100{index}", subsidiary_name="بانک", normalized_name="بانک", closing_debit=amount, closing_credit=0, net_closing_balance=amount, source_row_number=index))
        session.add(AccountMappingRule(organization_id=None, name="نقد", source_level="subsidiary", source_code_prefix="1110", target_statement="BS", target_field_id="BS.CASH", amount_source="net_closing", sign_multiplier=1, is_global=True, created_by="user-1"))
        session.flush(); AccountMappingService(session).apply(import_); session.flush()
        decisions = list(session.scalars(select(AccountMappingDecision)))
        assert all(item.status == "mapped" for item in decisions)
        first = FinancialStatementCalculationService(session).calculate(import_, "user-1")
        session.flush()
        value = session.scalar(select(FinancialStatementValue).where(FinancialStatementValue.run_id == first.id, FinancialStatementValue.field_id == "BS.CASH"))
        assert value.final_value == Decimal("60000000")
        second = FinancialStatementCalculationService(session).calculate(import_, "user-1")
        assert (first.version, second.version) == (1, 2)
        assert convert_money(Decimal("10"), "toman", "rial") == 100


def test_complete_workspace_inputs_and_money_units():
    organization = OrganizationCreate(
        entity_type="company",
        name="شرکت نمونه",
        national_id="10101234567",
        economic_code="411111111111",
        registration_number="12345",
        tax_file_number="987654",
        province="تهران",
        city="تهران",
        postal_code="1234567890",
        address="تهران، خیابان نمونه",
    )
    assert organization.entity_type == "company"
    assert organization.postal_code == "1234567890"

    fiscal_year = FiscalYearCreate(
        organization_id="organization-1",
        title="سال مالی ۱۴۰۵",
        start_date="2026-03-21",
        end_date="2027-03-20",
        status="open",
    )
    assert fiscal_year.status == "open"

    amount = Decimal("1000000")
    assert convert_money(amount, "rial", "rial") == amount
    assert convert_money(amount, "toman", "rial") == Decimal("10000000")


def test_exports_are_valid_files():
    field = FinancialStatementField(id="BS.CASH", statement="BS", title="موجودی نقد", field_type="mapped")
    value = FinancialStatementValue(field_id=field.id, final_value=Decimal("123456"), status="calculated")
    excel = workbook_export([(field, value)], "شرکت آزمون", "۱۴۰۵")
    pdf = pdf_export([(field, value)], "شرکت آزمون", "۱۴۰۵")
    assert excel.startswith(b"PK")
    assert pdf.startswith(b"%PDF")
    assert len(excel) > 1_000
    assert len(pdf) > 10_000


def test_parent_child_accounts_are_not_double_counted():
    def row(id_, general, subsidiary="", detail=""):
        return type("Row", (), {
            "id": id_, "general_code": general, "general_name": general,
            "subsidiary_code": subsidiary, "subsidiary_name": subsidiary,
            "detail_code": detail, "detail_name": detail,
        })()

    rows = [
        row("general", "1110"),
        row("subsidiary", "1110", "111001"),
        row("detail-1", "1110", "111001", "person-1"),
        row("detail-2", "1110", "111001", "person-2"),
        row("other-subsidiary", "1110", "111002"),
    ]
    assert FinancialStatementCalculationService._eligible_row_ids(rows) == {
        "detail-1", "detail-2", "other-subsidiary",
    }


def test_equity_mapping_is_specific_and_not_generic_311():
    assert AccountMappingService._fallback("311001", "سرمایه") == ("BS.CAPITAL", -1)
    assert AccountMappingService._fallback("311002", "سرمایه در جریان") == ("BS.CAPITAL_IN_PROGRESS", -1)
    assert AccountMappingService._fallback("311101", "اندوخته قانونی") == ("BS.LEGAL_RESERVE", -1)
    assert AccountMappingService._fallback("311201", "سود انباشته") == ("BS.RETAINED_EARNINGS", -1)
    assert AccountMappingService._fallback("311301", "سود جاری") == ("BS.CURRENT_YEAR_PROFIT_LOSS", -1)


def test_excel_and_pdf_use_the_same_export_dataset():
    rows = []
    for index, (field_id, title, amount) in enumerate((
        ("BS.CASH", "موجودی نقد", Decimal("123456")),
        ("BS.TOTAL_ASSETS", "جمع دارایی‌ها", Decimal("987654")),
        ("PL.NET_PROFIT", "سود خالص", Decimal("514198867")),
    )):
        statement = field_id.split(".", 1)[0]
        field = FinancialStatementField(id=field_id, statement=statement, title=title, field_type="mapped", sort_order=index)
        value = FinancialStatementValue(field_id=field_id, final_value=amount, status="calculated")
        rows.append((field, value))
    dto = build_export_data(rows, "شرکت آزمون", "۱۴۰۵", "rial")
    excel = workbook_export(rows, dto.company, dto.year, dto.money_unit)
    pdf = pdf_export(rows, dto.company, dto.year, dto.money_unit)
    workbook = load_workbook(BytesIO(excel), data_only=False)
    assert workbook["صورت وضعیت مالی"]["B5"].value == Decimal("123456")
    assert workbook["صورت وضعیت مالی"]["B6"].value == Decimal("987654")
    assert workbook["صورت سود و زیان"]["B5"].value == Decimal("514198867")
    document = fitz.open(stream=pdf, filetype="pdf")
    pdf_text = "\n".join(page.get_text() for page in document).replace(",", "")
    for amount in ("123456", "987654", "514198867"):
        assert amount in pdf_text


def test_current_profit_control_account_is_excluded_when_balance_sheet_already_balances(tmp_path):
    engine = create_engine(f"sqlite+pysqlite:///{tmp_path / 'profit-reconciliation.db'}")
    Base.metadata.create_all(engine)
    with Session(engine) as session:
        fields = {
            "PL.NET_PROFIT": ("PL", Decimal("514198867")),
            "BS.CURRENT_YEAR_PROFIT_LOSS": ("BS", Decimal("-514198867")),
            "BS.TOTAL_ASSETS": ("BS", Decimal("1259160010991")),
            "BS.TOTAL_LIABILITIES": ("BS", Decimal("353437593707")),
            "BS.CAPITAL": ("BS", Decimal("1000000")),
            "BS.CAPITAL_IN_PROGRESS": ("BS", Decimal("680000000000")),
            "BS.LEGAL_RESERVE": ("BS", Decimal("7698437671")),
            "BS.RETAINED_EARNINGS": ("BS", Decimal("218022979613")),
        }
        for index, (field_id, (statement, amount)) in enumerate(fields.items()):
            session.add(FinancialStatementField(id=field_id, statement=statement, title=field_id, field_type="mapped", sort_order=index))
            session.add(FinancialStatementValue(run_id="run-1", field_id=field_id, calculated_value=amount, final_value=amount, status="calculated"))
        session.flush()
        decision = type("Decision", (), {"target_field_id": "PL.OPERATING_REVENUE"})()
        status, difference, _ = FinancialStatementCalculationService(session)._reconcile_current_profit_loss(
            "run-1", [(decision, object())], Decimal(0)
        )
        current = session.scalar(select(FinancialStatementValue).where(
            FinancialStatementValue.run_id == "run-1",
            FinancialStatementValue.field_id == "BS.CURRENT_YEAR_PROFIT_LOSS",
        ))
        assert status == "passed"
        assert difference == 0
        assert current.final_value == 0
        assert current.status == "control_account_excluded"
